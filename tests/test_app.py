import io
import json
import unittest
import zipfile
import docx
from pypdf import PdfWriter

from app import create_app
from config import TestingConfig
from core.extractor import extract_text_from_file
from core.ai_detector import AIDetector
from core.citation_validator import CitationValidator
from core.vector_engine import VectorSearchEngine
from core.report_generator import ReportGenerator
from core.sanitizer import TextSanitizer
from core.checker import PlagiarismChecker
from core.phd_auditor import PhdResearchAuditor
from core.student_coach import AcademicStudentCoach


class EnterpriseAcademicOriginalityTestSuite(unittest.TestCase):
    def setUp(self):
        self.app_instance = create_app(TestingConfig)
        self.client = self.app_instance.test_client()
        self.checker = self.app_instance.checker

    def test_extract_plain_text(self):
        content = b"This is a plain text sample for unit testing."
        stream = io.BytesIO(content)
        extracted = extract_text_from_file(stream, "sample.txt")
        self.assertEqual(extracted, "This is a plain text sample for unit testing.")

    def test_extract_docx(self):
        doc = docx.Document()
        doc.add_paragraph("Artificial intelligence and natural language processing.")
        doc.add_paragraph("Second paragraph in docx test.")
        stream = io.BytesIO()
        doc.save(stream)
        stream.seek(0)

        extracted = extract_text_from_file(stream, "test.docx")
        self.assertIn("Artificial intelligence and natural language processing.", extracted)
        self.assertIn("Second paragraph in docx test.", extracted)

    def test_extract_pdf(self):
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        stream = io.BytesIO()
        writer.write(stream)
        stream.seek(0)

        extracted = extract_text_from_file(stream, "test.pdf")
        self.assertIsInstance(extracted, str)

    def test_ai_detector(self):
        detector = AIDetector()
        sample_text = (
            "Furthermore, it is important to note that the holistic approach plays a pivotal role in the seamless integration. "
            "In conclusion, this vibrant tapestry is a testament to the realm of multifaceted machine learning. "
            "Moreover, it is worth noting that deep neural networks foster innovation."
        )
        res = detector.analyze(sample_text)
        self.assertIn("ai_probability", res)
        self.assertIn("burstiness", res)
        self.assertGreater(res["ai_probability"], 0.0)

    def test_citation_validator(self):
        validator = CitationValidator()
        sample_text = (
            "Machine learning has revolutionized inferential statistics (Smith et al., 2024). "
            "According to Johnson (2023), deep neural networks generalize patterns across domains.\n\n"
            "References\n"
            "Smith, J., Doe, A., & Lee, K. (2024). Deep Learning Advances. Journal of AI, 12(3), 45-56.\n"
            "Johnson, R. (2023). Statistical Generalization in Neural Networks. Springer."
        )
        res = validator.validate_citations(sample_text)
        self.assertTrue(res["has_bibliography"])
        self.assertEqual(res["bibliography_entries_count"], 2)
        self.assertGreater(res["in_text_citations_count"], 0)

    def test_vector_engine(self):
        engine = VectorSearchEngine(vector_dim=128)
        engine.add_document("doc1", "Supervised learning models predict continuous outcomes via linear regression.")
        engine.add_document("doc2", "Quantum physics describes the behavior of matter and light on an atomic scale.")

        results = engine.search_top_k("Linear regression predicting continuous variables in machine learning.", k=1)
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]["doc_id"], "doc1")
        self.assertGreater(results[0]["similarity"], 0.0)

    def test_cosine_similarity_exact_match(self):
        sample = "This is a document for the Simple Plagiarism Checker. Hope you like it!"
        res = self.checker.analyze(sample, include_web_sources=False)
        self.assertGreater(res["overall_similarity"], 80.0)
        self.assertIn("database1.txt", [s["filename"] for s in res["sources_breakdown"]])

    def test_api_check_json_with_ai_and_diff(self):
        res = self.client.post("/check", json={
            "query": "Cybersecurity is the practice of protecting systems and networks from digital attacks.",
            "include_web": False,
            "private_draft": True
        })
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertIn("overall_similarity", data)
        self.assertIn("safeassign_risk", data)
        self.assertIn("ai_analysis", data)
        self.assertIn("diff_matches", data)
        self.assertIn("readability", data)
        self.assertIn("obfuscation_info", data)
        self.assertTrue(data.get("is_private_draft"))

    def test_text_sanitizer_obfuscation(self):
        # Insert Cyrillic 'a' (\u0430) and zero-width space (\u200b)
        malicious = "Th\u0430is text h\u200bas hidden Cyrillic homoglyphs."
        res = TextSanitizer.analyze_and_sanitize(malicious)
        self.assertTrue(res["has_obfuscation"])
        self.assertEqual(res["homoglyphs_detected"], 1)
        self.assertEqual(res["zero_width_chars_detected"], 1)
        self.assertEqual(res["sanitized_text"], "Thais text has hidden Cyrillic homoglyphs.")

    def test_readability_metrics(self):
        text = "Deep neural networks are composed of multiple layers to learn representations of data with multiple levels of abstraction."
        res = TextSanitizer.compute_readability(text)
        self.assertIn("flesch_reading_ease", res)
        self.assertIn("grade_level", res)
        self.assertGreater(res["avg_sentence_len"], 0.0)

    def test_smart_citations_generation(self):
        citations = PlagiarismChecker.generate_smart_citations(
            "machine_learning_foundations.txt", "https://arxiv.org/abs/2304.12345"
        )
        self.assertIn("apa", citations)
        self.assertIn("mla", citations)
        self.assertIn("ieee", citations)
        self.assertIn("Machine Learning Foundations", citations["apa"])

    def test_batch_zip_processing(self):
        # Create a test ZIP in memory with 2 student submissions
        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, 'w') as zf:
            zf.writestr("student1_essay.txt", "Cybersecurity is the practice of protecting systems from digital attacks.")
            zf.writestr("student2_essay.txt", "Zebras walk across violet fields on Neptune drinking cold tea.")
        zip_buffer.seek(0)

        res = self.client.post("/check/batch", data={
            "files": (zip_buffer, "class_test.zip"),
            "include_web": "false"
        }, content_type="multipart/form-data")

        self.assertEqual(res.status_code, 200)
        batch_data = res.get_json()
        self.assertEqual(batch_data["total_submissions"], 2)
        self.assertEqual(len(batch_data["gradebook"]), 2)

    def test_report_generator_endpoint(self):
        sample_payload = {
            "overall_similarity": 45.2,
            "safeassign_risk": "Medium Risk",
            "status_class": "warning",
            "total_words": 150,
            "highlighted_sentences": [
                {"text": "Sample flagged sentence.", "is_plagiarized": True, "source": "Sample Source", "similarity": 85.0}
            ],
            "sources_breakdown": [
                {"filename": "Sample Source", "similarity": 85.0, "source_word_count": 200, "badge": "📁 Local"}
            ]
        }
        res = self.client.post("/reports/html", json=sample_payload)
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"ScaleSynthAI SafeAssign Originality Report", res.data)

    def test_student_certificate_endpoint(self):
        sample_payload = {
            "student_name": "Jane Doe",
            "paper_title": "Ethics in Artificial Intelligence",
            "data": {
                "overall_similarity": 8.4,
                "total_words": 1250,
                "ai_analysis": {"ai_probability": 12.0},
                "citation_analysis": {"in_text_citations_count": 14},
                "readability": {"grade_level": "Grade 12.1"}
            }
        }
        res = self.client.post("/reports/certificate", json=sample_payload)
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"Certificate of Academic Authorship", res.data)
        self.assertIn(b"Jane Doe", res.data)
        self.assertIn(b"Ethics in Artificial Intelligence", res.data)

    def test_openapi_spec_and_docs(self):
        res_spec = self.client.get("/api/spec.json")
        self.assertEqual(res_spec.status_code, 200)
        spec_data = res_spec.get_json()
        self.assertEqual(spec_data["openapi"], "3.0.3")
        self.assertIn("/check", spec_data["paths"])

    def test_extract_latex(self):
        latex_content = rb"""
        \documentclass{article}
        % This is a LaTeX comment that should be ignored
        \begin{document}
        \section{Introduction}
        Deep learning methods achieve state of the art results \cite{vaswani2017attention}.
        Here is an equation:
        \begin{equation}
        \mathcal{L}(\theta) = -\log p(x)
        \end{equation}
        \end{document}
        """
        stream = io.BytesIO(latex_content)
        extracted = extract_text_from_file(stream, "paper.tex")
        self.assertIn("Deep learning methods achieve state of the art results", extracted)
        self.assertIn("(vaswani2017attention, 2024)", extracted)
        self.assertNotIn("This is a LaTeX comment", extracted)

    def test_extract_jupyter_notebook(self):
        nb_json = {
            "cells": [
                {"cell_type": "markdown", "source": ["# Graph Neural Networks\n", "This research explores message passing."]},
                {"cell_type": "code", "source": ["def train():\n", "    '''Executes training step'''\n", "    pass"]}
            ]
        }
        stream = io.BytesIO(json.dumps(nb_json).encode('utf-8'))
        extracted = extract_text_from_file(stream, "experiment.ipynb")
        self.assertIn("Graph Neural Networks", extracted)
        self.assertIn("Executes training step", extracted)

    def test_phd_auditor_anonymity_violation(self):
        sample = (
            "In our previous paper [Smith et al.], we introduced model X. "
            "Our source code is publicly accessible at https://github.com/myusername/my-secret-model."
        )
        res = PhdResearchAuditor.audit_manuscript(sample)
        self.assertFalse(res["is_anonymity_compliant"])
        self.assertGreaterEqual(res["anonymity_count"], 2)

    def test_delete_source_with_pin(self):
        # 1. Upload a dummy source first
        stream = io.BytesIO(b"Institutional document content for deletion test.")
        upload_res = self.client.post("/sources/upload", data={
            "file": (stream, "temp_delete_test.txt")
        }, content_type="multipart/form-data")
        self.assertEqual(upload_res.status_code, 201)

        # 2. Attempt delete without PIN -> 403 Forbidden
        del_no_pin = self.client.delete("/sources/temp_delete_test.txt")
        self.assertEqual(del_no_pin.status_code, 403)
        self.assertIn("Invalid or missing Admin PIN", del_no_pin.get_json()["error"])

        # 3. Attempt delete with wrong PIN -> 403 Forbidden
        del_wrong_pin = self.client.delete("/sources/temp_delete_test.txt", headers={"X-Admin-PIN": "wrong_pin"})
        self.assertEqual(del_wrong_pin.status_code, 403)

        # 4. Attempt delete with correct PIN -> 200 OK
        del_correct_pin = self.client.delete(
            "/sources/temp_delete_test.txt",
            headers={"X-Admin-PIN": self.app_instance.config.get("ADMIN_PIN", "1234")}
        )
        self.assertEqual(del_correct_pin.status_code, 200)
        self.assertTrue(del_correct_pin.get_json()["success"])

    def test_false_positive_noun_collision_defense(self):
        # A scientific text that has domain words ('price', 'seaborn', 'diamond') but zero copied sentences
        sample = (
            "We analyzed diamond price distribution using the Python seaborn package. "
            "The skewness coefficient was 1.62 with a Pearson correlation r = 0.92 across clarity categories. "
            "Our statistical tests demonstrate that logarithmic transformation stabilizes variance."
        )
        res = self.checker.analyze(sample, include_web_sources=False)
        self.assertEqual(res["overall_similarity"], 0.0)
        self.assertEqual(res["safeassign_risk"], "Low Risk")
        self.assertEqual(len(res["diff_matches"]), 0)

    def test_academic_paraphraser_endpoint(self):
        res = self.client.post("/api/paraphrase", json={
            "sentence": "Deep learning models predict continuous outcomes via linear regression.",
            "source_name": "Machine Learning Foundations.txt",
            "source_title": "Machine Learning Foundations"
        })
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertIn("suggestions", data)
        self.assertEqual(len(data["suggestions"]), 3)
        self.assertIn("Machine Learning Foundations", data["suggestions"][0]["text"])

    def test_citation_generator_endpoint(self):
        # Test arXiv citation generation
        res = self.client.post("/api/cite", json={
            "query": "1706.03762"
        })
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertIn("bibtex", data)
        self.assertIn("apa", data)
        self.assertIn("mla", data)
        self.assertIn("ieee", data)
        self.assertIn("@article", data["bibtex"])

    def test_draft_comparator_endpoint(self):
        v1 = "Linear regression models predict outcomes. Gradient descent optimizes parameters."
        v2 = "Linear regression models predict continuous outcomes. Adam optimizer accelerates convergence. Evaluation metrics confirm robustness."
        res = self.client.post("/check/compare-drafts", json={
            "draft_v1": v1,
            "draft_v2": v2
        })
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertIn("similarity_percentage", data)
        self.assertIn("revision_percentage", data)
        self.assertIn("added_count", data)
        self.assertGreater(data["draft_v2_words"], data["draft_v1_words"])

    def test_student_coach_claims_and_tone(self):
        # Unsupported claim without citation vs supported claim
        text = (
            "Recent studies show that transformer networks optimize representation learning. "
            "However, 85% of deep models require parameter pruning. "
            "According to Vaswani et al. (2017), self-attention computes dense pairwise representations. "
            "We basically looked at a lot of parameters to get rid of big impacts."
        )
        unsupported = AcademicStudentCoach.scan_unsupported_claims(text)
        self.assertGreaterEqual(len(unsupported), 1)
        # Should flag "Recent studies show" or "85% of", but NOT Vaswani (which has citation)
        sentences_flagged = [u["sentence"] for u in unsupported]
        self.assertTrue(any("studies show" in s.lower() or "85%" in s for s in sentences_flagged))
        self.assertFalse(any("Vaswani" in s for s in sentences_flagged))

        # Tone suggestions
        tone_suggestions = AcademicStudentCoach.analyze_tone_and_vocabulary(text)
        self.assertGreater(len(tone_suggestions), 0)
        matched = [t["matched_term"].lower() for t in tone_suggestions]
        self.assertTrue(any("basically" in m or "a lot of" in m or "get rid of" in m for m in matched))

    def test_student_coach_alphabetizer(self):
        raw_refs = (
            "Vaswani, A. et al. (2017). Attention is all you need.\n"
            "Brown, T. et al. (2020). Language models are few-shot learners. https://doi.org/10.1234/test\n"
            "Devlin, J. et al. (2018). BERT: Pre-training of transformers.\n"
            "Achiam, J. et al. GPT-4 technical report."
        )
        res = AcademicStudentCoach.alphabetize_and_format_references(raw_refs)
        self.assertEqual(res["count"], 4)
        sorted_refs = res["sorted_references"]
        # Alphabetical order: Achiam, Brown, Devlin, Vaswani
        self.assertTrue(sorted_refs[0].startswith("Achiam"))
        self.assertTrue(sorted_refs[1].startswith("Brown"))
        self.assertTrue(sorted_refs[2].startswith("Devlin"))
        self.assertTrue(sorted_refs[3].startswith("Vaswani"))
        # Issue flagged for Achiam (missing year)
        self.assertGreater(len(res["issues"]), 0)

    def test_student_coach_thesis_evaluator(self):
        good_abstract = (
            "In this paper, we hypothesize that self-supervised pre-training improves inductive bias. "
            "We evaluate our framework using empirical datasets across 5 benchmark domains. "
            "Our results demonstrate fundamental significance for trustworthy natural language processing."
        )
        res = AcademicStudentCoach.evaluate_thesis_abstract(good_abstract)
        self.assertGreaterEqual(res["score"], 80)
        self.assertTrue(res["has_hypothesis"])
        self.assertTrue(res["has_method"])
        self.assertTrue(res["has_significance"])

    def test_student_coach_endpoints(self):
        # 1. Tone and Claims endpoint
        res1 = self.client.post("/api/student-coach/tone-and-claims", json={
            "text": "Studies show that machine learning is basically a big deal."
        })
        self.assertEqual(res1.status_code, 200)
        data1 = res1.get_json()
        self.assertIn("unsupported_claims", data1)
        self.assertIn("tone_suggestions", data1)

        # 2. Alphabetize endpoint
        res2 = self.client.post("/api/student-coach/alphabetize-references", json={
            "references": "Zeta, A. (2020). Test.\nAlpha, B. (2021). Test."
        })
        self.assertEqual(res2.status_code, 200)
        data2 = res2.get_json()
        self.assertEqual(data2["sorted_references"][0][:5], "Alpha")

        # 3. Evaluate thesis endpoint
        res3 = self.client.post("/api/student-coach/evaluate-thesis", json={
            "text": "We propose a novel method for anomaly detection using graph neural networks."
        })
        self.assertEqual(res3.status_code, 200)
        data3 = res3.get_json()
        self.assertIn("score", data3)

    def test_check_includes_student_coach_data(self):
        res = self.client.post("/check", json={
            "query": "Studies show that deep learning improves perception. We basically evaluated 500 samples.",
            "include_web": False,
            "private_draft": True
        })
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertIn("student_coach", data)
        self.assertIn("unsupported_claims", data["student_coach"])
        self.assertIn("tone_suggestions", data["student_coach"])
        self.assertIn("thesis_evaluation", data["student_coach"])


if __name__ == "__main__":
    unittest.main()



